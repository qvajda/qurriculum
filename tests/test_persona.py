import docx
import yaml

from qurriculum import generate, persona
from test_personas import extract_text, walk_values

SAMPLE = persona.DEFAULT_OUT.parent / "camille_dubois.yml"
CONVENTION_HEADINGS = [
    "Formation",
    "Expérience professionnelle",
    "Langues",
    "Compétences informatiques",
    "Divers",
]  # conclusion SECTIONS, coordonnées being the header block


def load():
    return yaml.safe_load(SAMPLE.read_text(encoding="utf-8"))


def test_persona_is_raw_facts():
    data = load()
    assert "sections" not in data
    assert not set(persona.HEADINGS.values()) & set(walk_values(data))


def test_headings_and_order_match_conventions():
    content = persona.build_content(load())
    assert [s["heading"] for s in content["sections"]] == CONVENTION_HEADINGS


def test_most_recent_first_and_date_format():
    content = persona.build_content(load())
    employers = content["sections"][1]["entries"]
    assert employers[0].startswith("09/2021 – présent : ")
    assert employers[1].startswith("09/2020 – 06/2021 : ")


def test_every_fact_reaches_docx_deterministically(tmp_path):
    data = load()
    outs = [generate.render(persona.build_content(data), tmp_path / f"{i}.docx") for i in "ab"]
    texts = [extract_text(docx.Document(o)) for o in outs]
    assert texts[0] == texts[1]
    body = "\n".join(texts[0])
    facts = [
        v
        for v in walk_values({k: v for k, v in data.items()})
        if not v[:4].isdigit() or v.startswith("+")
    ]
    assert facts and all(f in body for f in facts), [f for f in facts if f not in body]


def test_stub_writer_is_used_no_model():
    content = persona.build_content(load(), write=lambda kind, fact: f"<{kind}>")
    assert content["sections"][3]["entries"][0] == "<skill>"


def test_committed_sample_is_generated():
    assert docx.Document(persona.DEFAULT_OUT).paragraphs
