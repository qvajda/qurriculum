from pathlib import Path

import docx
import pytest
import yaml

from qurriculum import generate

PERSONAS = sorted(Path(__file__).parent.joinpath("personas").glob("*.yml"))


def extract_text(document: docx.Document) -> list[str]:
    text = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text.extend(p.text for p in cell.paragraphs)
    return text


def walk_values(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for v in value.values():
            yield from walk_values(v)
    elif isinstance(value, list):
        for v in value:
            yield from walk_values(v)


def test_at_least_four_personas():
    assert len(PERSONAS) >= 4


@pytest.mark.parametrize("persona_path", PERSONAS, ids=lambda p: p.stem)
def test_persona_injects_without_shattering(persona_path, tmp_path):
    expected = [s["heading"] for s in generate.load_content()["sections"]]
    persona = yaml.safe_load(persona_path.read_text(encoding="utf-8"))
    out = tmp_path / f"{persona_path.stem}.docx"
    generate.render(persona, out)

    text = extract_text(docx.Document(out))
    full_text = "\n".join(text)

    assert [t for t in text if t in expected] == expected

    for value in walk_values(persona):
        assert value in full_text, (persona_path.stem, value)
